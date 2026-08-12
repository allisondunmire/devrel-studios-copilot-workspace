import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Create workbook
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Filming Schedule"

# Define headers
headers = ["Active Projects", "EP", "Type", "Relevant Dates", "Filming Location", "Status", "Action Items/Questions"]
ws.append(headers)

# Style header row
header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF", size=11)
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin_border

# Project data from ADO
projects = [
    ["Fabric Tech Talk Fridays", "Golnaz Alibeigi", "Show", "", "Studio", "", ""],
    ["Cozy AI Kitchen", "Ross Heise", "Show", "", "Studio", "", ""],
    ["Sip and Sync with Azure", "Ross Heise", "Show", "", "", "", ""],
    ["Open at Microsoft", "Kaitlin McKinnon", "Show", "", "Streamyard (self service)", "", ""],
    ["The Low Code Revolution", "Cameron Tomisser", "Show", "", "", "", ""],
    ["Data Exposed", "Ross Heise", "Show", "", "", "", ""],
    ["Microsoft MVP Unplugged", "Sara Finkelstein", "Show", "", "", "", ""],
    ["Azure Friday 🌮", "Kaitlin McKinnon", "Show", "", "Streamyard (self service)", "", ""],
    ["Quest to Compile", "Sara Finkelstein", "Show", "", "Streamyard (self service)", "", ""],
    ["Fabric Executive Insights", "Golnaz Alibeigi", "Show", "", "", "", ""],
]

# Add project data
for row_idx, project in enumerate(projects, start=2):
    for col_idx, value in enumerate(project, start=1):
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.value = value
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        cell.border = thin_border

# Set column widths
column_widths = [30, 20, 12, 20, 25, 15, 30]
for idx, width in enumerate(column_widths, start=1):
    ws.column_dimensions[get_column_letter(idx)].width = width

# Set row height for header
ws.row_dimensions[1].height = 30

# Freeze top row
ws.freeze_panes = "A2"

# Save file
file_path = "c:\\Users\\v-adunmire\\WORK\\devrel-studios-copilot-workspace\\DevRel_Studios_Filming_Tracker.xlsx"
wb.save(file_path)
print(f"✓ Excel tracker created: {file_path}")
print(f"✓ {len(projects)} active projects populated")
print("✓ Ready for you to add filming dates from calendar")
