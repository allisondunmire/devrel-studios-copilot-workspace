from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font

WORKBOOK_PATH = Path("DevRel_BirdsEye_WIP.xlsx")
ADO_BASE = "https://dev.azure.com/devrel/Studios/_workitems/edit/{}"
ADO_RE = re.compile(r"ADO\s*(\d+)", re.IGNORECASE)


def header_is_active_projects(value: object) -> bool:
    if not isinstance(value, str):
        return False
    normalized = " ".join(value.lower().split())
    return (
        "active project" in normalized
        or normalized == "projects"
        or normalized == "project"
    )


wb = load_workbook(WORKBOOK_PATH)
updated = 0

for ws in wb.worksheets:
    header_row = None
    for row in range(1, min(ws.max_row, 15) + 1):
        if header_is_active_projects(ws.cell(row=row, column=2).value):
            header_row = row
            break

    if header_row is None:
        continue

    for row in range(header_row + 1, ws.max_row + 1):
        cell = ws.cell(row=row, column=2)
        if not isinstance(cell.value, str):
            continue
        match = ADO_RE.search(cell.value)
        if not match:
            continue

        work_item_id = match.group(1)
        cell.hyperlink = ADO_BASE.format(work_item_id)
        cell.style = "Hyperlink"
        # Keep title text readable and visibly link-like across Office themes.
        cell.font = Font(color="0563C1", underline="single")
        updated += 1

wb.save(WORKBOOK_PATH)
print(f"Updated hyperlinks: {updated}")
